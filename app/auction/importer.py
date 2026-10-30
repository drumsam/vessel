"""Small, operator-authorized Christie's batch import into existing evidence.

No network, taxonomy or merchant pricing path. Unresolved rows use offer_reject.
"""

import csv
import hashlib
import json
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

from app.api.schemas import AuctionEvidenceIn
from app.auction import repo as auction_repo
from app.auction.evidence import InvalidAuctionEvidence
from app.db import repo

SCHEMA_VERSION = "auction-import-1"
SOURCE = "christies"


class ImportConflict(ValueError):
    pass


def file_rows(filename, content):
    if len(content) > 2_000_000:
        raise ValueError("File exceeds 2 MB")
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        reader = csv.DictReader(StringIO(content.decode("utf-8-sig")))
        if not reader.fieldnames or len(set(reader.fieldnames)) != len(
            reader.fieldnames
        ):
            raise ValueError("Unique column headers required")
        yield from reader
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 20_000_000:
                raise ValueError("Expanded workbook exceeds 20 MB")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            reader = book.active.iter_rows(values_only=True)
            headers = next(reader, ())
            if (
                not headers
                or any(not isinstance(h, str) for h in headers)
                or len(set(headers)) != len(headers)
            ):
                raise ValueError("Unique text headers required")
            for values in reader:
                if any(v is not None for v in values):
                    yield dict(zip(headers, values))
        finally:
            book.close()
    else:
        raise ValueError("Use CSV or XLSX")


def clean_row(row):
    return {
        str(k): v.isoformat() if isinstance(v, (date, datetime)) else v
        for k, v in row.items()
    }


def parse_row(row):
    # No quantity/format defaults: the lot's commercial facts must be stated.
    data = {
        k: row[k]
        for k in AuctionEvidenceIn.model_fields
        if row.get(k) not in (None, "")
    }
    for key in ("amount", "premium_rate", "fx_rate"):
        if key in data:
            data[key] = str(data[key])
    body = AuctionEvidenceIn.model_validate(data)
    if body.auction_house != "CHRISTIES":
        raise ValueError("This milestone supports the existing Christie's source only")
    vintage = int(str(row["vintage"]))
    bottle_ml = int(str(row["bottle_ml"]))
    if not (vintage == 1000 or 1800 <= vintage <= date.today().year) or bottle_ml <= 0:
        raise ValueError("Explicit vintage and positive bottle_ml required")
    if body.sale_date > datetime.now(UTC).date():
        raise ValueError("Future sale date")
    producer, wine = row.get("producer"), row.get("wine")
    if (
        not isinstance(producer, str)
        or not isinstance(wine, str)
        or not producer.strip()
        or not wine.strip()
    ):
        raise ValueError("Separate producer and wine fields required")
    return body, vintage, bottle_ml


def record_row(conn, envelope, *, operator_id):
    row = envelope["row"]
    body, vintage, bottle_ml = parse_row(row)
    registry = repo.load_registry(conn)
    wine_id = registry.resolve_candidate(
        row["producer"],
        row["wine"],
        row.get("lwin7") or None,
        row.get("release_code") or None,
    )
    if wine_id is None:
        return None
    fields = body.model_dump()
    event = ["CHRISTIES", body.sale_date.isoformat(), body.lot_ref.strip()]
    key = hashlib.sha256(json.dumps(event).encode()).hexdigest()
    # Serialize even concurrent imports of the same source event, before checking.
    conn.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))", (key,))
    prior = conn.execute(
        """SELECT * FROM auction_evidence WHERE retracted_at IS NULL AND
        (import_key=%s OR (auction_house=%s AND sale_date=%s AND lot_ref=%s))
        ORDER BY entered_at LIMIT 1""",
        (key, body.auction_house, body.sale_date, body.lot_ref),
    ).fetchone()
    facts = dict(fields, wine_id=wine_id, vintage=vintage, bottle_ml=bottle_ml)
    if prior:
        if any(
            (str(prior[k]) != str(v) if k == "wine_id" else prior[k] != v)
            for k, v in facts.items()
        ):
            raise ImportConflict(
                "Conflicting facts for this sale/lot; review and retract the prior result first"
            )
        return "duplicate"
    result = auction_repo.record_auction_evidence(
        conn, **facts, entered_by=operator_id, raw_evidence=envelope, import_key=key
    )
    return str(result["auction_evidence_id"])


def import_file(conn, *, filename, content, operator_id):
    rows = []
    for row in file_rows(filename, content):
        if len(rows) == 2000:
            raise ValueError("Maximum 2000 rows per batch")
        rows.append(row)
    when = datetime.now(UTC)
    snapshot = repo.create_snapshot(
        conn, captured_at=when, schema_version=SCHEMA_VERSION, is_demo=False
    )
    written = rejected = unresolved = duplicates = 0
    errors = []
    digest = hashlib.sha256(content).hexdigest()
    for number, raw in enumerate(rows, 2):
        envelope = {
            "kind": SCHEMA_VERSION,
            "filename": Path(filename).name,
            "file_sha256": digest,
            "row_number": number,
            "row": clean_row(raw),
            "operator_id": str(operator_id),
        }
        try:
            with conn.transaction():
                result = record_row(conn, envelope, operator_id=operator_id)
            if result is None:
                # Retain a distinct unresolved row once, using the existing reject store.
                existing = conn.execute(
                    """SELECT 1 FROM offer_reject WHERE merchant_id=%s
                    AND reason='UNRESOLVED_IDENTITY' AND raw->>'kind'=%s
                    AND raw->>'file_sha256'=%s AND raw->>'row_number'=%s LIMIT 1""",
                    (SOURCE, SCHEMA_VERSION, digest, str(number)),
                ).fetchone()
                if not existing:
                    repo.write_reject(
                        conn,
                        snapshot,
                        SOURCE,
                        "UNRESOLVED_IDENTITY",
                        "Auction import requires canonical resolution",
                        envelope,
                    )
                unresolved += 1
            elif result == "duplicate":
                duplicates += 1
            else:
                written += 1
        except (ValueError, KeyError, TypeError, InvalidAuctionEvidence) as exc:
            rejected += 1
            # Original invalid evidence remains auditable; no raw input in the response.
            repo.write_reject(
                conn,
                snapshot,
                SOURCE,
                "AUCTION_IMPORT_INVALID",
                type(exc).__name__,
                envelope,
            )
            errors.append(
                {
                    "row": number,
                    "error": "conflicting event; review prior result"
                    if isinstance(exc, ImportConflict)
                    else "invalid row; inspect retained original evidence",
                }
            )
    repo.write_snapshot_merchant(
        conn,
        snapshot,
        SOURCE,
        captured_at=when,
        scope_params={"kind": SCHEMA_VERSION},
        written=written,
        rejected=rejected + unresolved,
    )
    repo.finish_snapshot(
        conn,
        snapshot,
        attempted=1,
        succeeded=1,
        written=written,
        rejected=rejected + unresolved,
        failures=errors,
    )
    return {
        "snapshot_id": str(snapshot),
        "written": written,
        "unresolved": unresolved,
        "rejected": rejected,
        "duplicates": duplicates,
        "errors": errors,
    }


def reresolve_import(conn, row):
    envelope = row["raw"]
    if envelope.get("kind") != SCHEMA_VERSION:
        return False
    try:
        with conn.transaction():
            result = record_row(conn, envelope, operator_id=envelope["operator_id"])
            if result:
                repo.delete_reject(conn, str(row["reject_id"]))
        return bool(result)
    except (ValueError, KeyError, TypeError, InvalidAuctionEvidence):
        return False  # Still requires review; original remains intact.
